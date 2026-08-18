#!/usr/bin/env python3
"""Сбор базы потенциальных клиентов ООО «СКЛАД» через официальный 2GIS Places API.

Почему API, а не парсинг страниц: HTML-скрапинг 2GIS и Яндекс.Карт нарушает их
пользовательское соглашение и упирается в капчу после первых сотен запросов.
Places API отдаёт те же данные легально и стабильно.

Ключ берётся из переменной окружения DGIS_API_KEY (или флага --key).

Примеры:
    # все сегменты по городу, результат в XLSX
    python3 leadgen_2gis.py --city Казань --out leads_kazan.xlsx

    # только под мини-экскаваторы
    python3 leadgen_2gis.py --city Казань --equipment "Мини-экскаваторы"

    # конкретные сегменты, с ограничением глубины
    python3 leadgen_2gis.py --city Казань --segment "Аренда спецтехники" --pages 5

    # посмотреть план запросов, ничего не тратя из квоты
    python3 leadgen_2gis.py --city Казань --dry-run
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import time
from dataclasses import dataclass, field, asdict

import requests

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from segments import SEGMENTS, find_segments, segments_for_equipment, all_equipment  # noqa: E402

API_URL = "https://catalog.api.2gis.com/3.0/items"
FIELDS = ",".join(
    [
        "items.contact_groups",
        "items.address",
        "items.point",
        "items.rubrics",
        "items.org",
        "items.external_content",
    ]
)
# Places API отдаёт максимум 10 записей на страницу и не пускает глубже ~5 страниц
# на один запрос. Поэтому охват набирается количеством узких запросов, а не глубиной.
PAGE_SIZE = 10
MAX_PAGES = 5
THROTTLE_SEC = 0.25


@dataclass
class Lead:
    company: str
    phones: list[str] = field(default_factory=list)
    website: str = ""
    email: str = ""
    address: str = ""
    city: str = ""
    segment: str = ""
    equipment: str = ""
    rubrics: str = ""
    branches: int = 1
    score: int = 0
    pitch: str = ""
    org_id: str = ""

    def as_row(self) -> dict:
        row = asdict(self)
        row["phones"] = ", ".join(self.phones)
        return row


def normalize_phone(raw: str) -> str:
    """Российский номер к виду +7XXXXXXXXXX. Нераспознанное возвращается как есть."""
    digits = re.sub(r"\D", "", raw or "")
    if len(digits) == 11 and digits[0] in "78":
        return "+7" + digits[1:]
    if len(digits) == 10:
        return "+7" + digits
    return raw.strip()


def extract_contacts(item: dict) -> tuple[list[str], str, str]:
    """Телефоны, сайт и e-mail из contact_groups организации."""
    phones: list[str] = []
    website = ""
    email = ""
    for group in item.get("contact_groups", []) or []:
        for contact in group.get("contacts", []) or []:
            ctype = contact.get("type")
            value = contact.get("value") or contact.get("text") or ""
            if not value:
                continue
            if ctype == "phone":
                phone = normalize_phone(value)
                if phone and phone not in phones:
                    phones.append(phone)
            elif ctype == "website" and not website:
                website = contact.get("url") or value
            elif ctype == "email" and not email:
                email = value
    return phones, website, email


def score_lead(lead: Lead, weight: int) -> int:
    """Приоритет обзвона: 0–100.

    Вес сегмента — основа. Дальше добавляем за то, что делает лид рабочим:
    наличие телефона (без него лид мёртв), сеть филиалов (больше техники),
    собственный сайт (компания живая, а не заброшенная карточка).
    """
    score = weight * 20  # 20 / 40 / 60
    if lead.phones:
        score += 20
    if len(lead.phones) > 1:
        score += 5
    if lead.branches > 1:
        score += min(lead.branches * 2, 10)
    if lead.website:
        score += 5
    return min(score, 100)


def search(key: str, query: str, city: str, pages: int, verbose: bool) -> list[dict]:
    """Постранично забирает организации по одному запросу."""
    collected: list[dict] = []
    full_query = f"{query} {city}".strip()
    for page in range(1, pages + 1):
        params = {
            "q": full_query,
            "key": key,
            "fields": FIELDS,
            "page": page,
            "page_size": PAGE_SIZE,
            "type": "branch",
        }
        try:
            resp = requests.get(API_URL, params=params, timeout=30)
        except requests.RequestException as exc:
            print(f"  ! сетевая ошибка на «{full_query}» стр.{page}: {exc}", file=sys.stderr)
            break

        if resp.status_code == 403:
            print("  ! 403 — ключ отклонён. Проверь DGIS_API_KEY и лимиты тарифа.", file=sys.stderr)
            break
        if resp.status_code == 429:
            print("  ! 429 — превышен лимит запросов, пауза 5 с", file=sys.stderr)
            time.sleep(5)
            continue
        if resp.status_code != 200:
            print(f"  ! HTTP {resp.status_code} на «{full_query}»", file=sys.stderr)
            break

        payload = resp.json()
        result = payload.get("result") or {}
        items = result.get("items") or []
        if not items:
            break
        collected.extend(items)

        total = result.get("total", 0)
        if verbose:
            print(f"    «{full_query}» стр.{page}: +{len(items)} (всего в 2GIS: {total})")
        if page * PAGE_SIZE >= total:
            break
        time.sleep(THROTTLE_SEC)
    return collected


def collect(key: str, city: str, chosen: list[dict], pages: int, verbose: bool) -> list[Lead]:
    """Обходит сегменты и сводит результаты в дедуплицированный список лидов."""
    by_id: dict[str, Lead] = {}

    for seg in chosen:
        print(f"[{seg['segment']}] запросов: {len(seg['queries'])}")
        for query in seg["queries"]:
            for item in search(key, query, city, pages, verbose):
                org_id = str((item.get("org") or {}).get("id") or item.get("id") or "")
                name = (item.get("name") or "").strip()
                if not name:
                    continue
                dedup_key = org_id or name.lower()

                if dedup_key in by_id:
                    # Компания уже найдена другим запросом: не плодим дубль, а
                    # расширяем её профиль потребности — значит она подходит под
                    # несколько сегментов сразу, это более горячий лид.
                    existing = by_id[dedup_key]
                    if seg["segment"] not in existing.segment:
                        existing.segment += f"; {seg['segment']}"
                        existing.score = min(existing.score + 5, 100)
                    for eq in seg["equipment"]:
                        if eq not in existing.equipment:
                            existing.equipment += f", {eq}"
                    continue

                phones, website, email = extract_contacts(item)
                rubrics = ", ".join(r.get("name", "") for r in (item.get("rubrics") or []))
                lead = Lead(
                    company=name,
                    phones=phones,
                    website=website,
                    email=email,
                    address=(item.get("address_name") or (item.get("address") or {}).get("name") or ""),
                    city=city,
                    segment=seg["segment"],
                    equipment=", ".join(seg["equipment"]),
                    rubrics=rubrics,
                    branches=int((item.get("org") or {}).get("branch_count") or 1),
                    pitch=seg["pitch"],
                    org_id=org_id,
                )
                lead.score = score_lead(lead, seg["weight"])
                by_id[dedup_key] = lead

    leads = list(by_id.values())
    leads.sort(key=lambda l: (-l.score, l.company))
    return leads


COLUMNS = [
    ("company", "Компания", 42),
    ("phones", "Телефон", 26),
    ("website", "Сайт", 26),
    ("email", "E-mail", 24),
    ("address", "Адрес", 40),
    ("segment", "Сегмент", 34),
    ("equipment", "Что предлагать", 40),
    ("score", "Приоритет", 11),
    ("branches", "Филиалов", 10),
    ("pitch", "Аргумент для звонка", 60),
    ("rubrics", "Рубрики 2GIS", 34),
]


def write_xlsx(leads: list[Lead], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col, (_, title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    hot = PatternFill("solid", fgColor="C6EFCE")
    warm = PatternFill("solid", fgColor="FFEB9C")

    for row_idx, lead in enumerate(leads, start=2):
        data = lead.as_row()
        for col, (field_name, _, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col, value=data.get(field_name, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=field_name in ("pitch", "equipment", "segment"))
        score_cell = ws.cell(row=row_idx, column=8)
        if lead.score >= 80:
            score_cell.fill = hot
        elif lead.score >= 60:
            score_cell.fill = warm

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(leads) + 1}"
    wb.save(path)


def write_csv(leads: list[Lead], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=[c[0] for c in COLUMNS], extrasaction="ignore")
        writer.writerow({c[0]: c[1] for c in COLUMNS})
        for lead in leads:
            writer.writerow(lead.as_row())


def main() -> int:
    parser = argparse.ArgumentParser(description="Сбор базы клиентов по потребности в спецтехнике (2GIS Places API)")
    parser.add_argument("--city", help="Город поиска, например: Казань")
    parser.add_argument("--segment", action="append", default=[], help="Сегмент (можно повторять). По умолчанию — все")
    parser.add_argument("--equipment", help="Отобрать сегменты под конкретную технику, например «Мини-экскаваторы»")
    parser.add_argument("--out", default="leads.xlsx", help="Файл результата (.xlsx или .csv)")
    parser.add_argument("--pages", type=int, default=MAX_PAGES, help=f"Страниц на запрос (по умолчанию {MAX_PAGES})")
    parser.add_argument("--key", default=os.environ.get("DGIS_API_KEY", ""), help="Ключ 2GIS (или env DGIS_API_KEY)")
    parser.add_argument("--min-score", type=int, default=0, help="Отбросить лиды с приоритетом ниже порога")
    parser.add_argument("--with-phone-only", action="store_true", help="Оставить только записи с телефоном")
    parser.add_argument("--dry-run", action="store_true", help="Показать план запросов и выйти")
    parser.add_argument("--list-segments", action="store_true", help="Показать сегменты и номенклатуру")
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args()

    if args.list_segments:
        print("Номенклатура:", ", ".join(all_equipment()), "\n")
        for seg in SEGMENTS:
            print(f"  [вес {seg['weight']}] {seg['segment']}")
            print(f"      техника: {', '.join(seg['equipment'])}")
            print(f"      запросы: {', '.join(seg['queries'])}")
        return 0

    if not args.city:
        parser.error("нужен --city (или --list-segments для просмотра сегментов)")

    if args.equipment:
        chosen = segments_for_equipment(args.equipment)
        if not chosen:
            print(f"Под «{args.equipment}» сегментов нет. Доступно: {', '.join(all_equipment())}", file=sys.stderr)
            return 1
    else:
        chosen = find_segments(args.segment)
        if not chosen:
            print(f"Сегменты не найдены: {args.segment}. Список: --list-segments", file=sys.stderr)
            return 1

    total_queries = sum(len(s["queries"]) for s in chosen)
    print(f"Город: {args.city}")
    print(f"Сегментов: {len(chosen)}, запросов: {total_queries}, до {total_queries * args.pages} обращений к API\n")

    if args.dry_run:
        for seg in chosen:
            print(f"[{seg['segment']}] → {', '.join(seg['equipment'])}")
            for q in seg["queries"]:
                print(f"    q = «{q} {args.city}»")
        return 0

    if not args.key:
        print("Нет ключа API. Задай DGIS_API_KEY или передай --key.", file=sys.stderr)
        print("Ключ выпускается в личном кабинете 2GIS Platform Manager.", file=sys.stderr)
        return 1

    leads = collect(args.key, args.city, chosen, args.pages, args.verbose)

    if args.with_phone_only:
        leads = [l for l in leads if l.phones]
    if args.min_score:
        leads = [l for l in leads if l.score >= args.min_score]

    if not leads:
        print("\nНичего не найдено. Проверь город и лимиты ключа.")
        return 0

    if args.out.endswith(".csv"):
        write_csv(leads, args.out)
    else:
        write_xlsx(leads, args.out)

    with_phone = sum(1 for l in leads if l.phones)
    hot = sum(1 for l in leads if l.score >= 80)
    print(f"\nСобрано компаний: {len(leads)}")
    print(f"  с телефоном:    {with_phone}")
    print(f"  приоритет 80+:  {hot}")
    print(f"  файл:           {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
